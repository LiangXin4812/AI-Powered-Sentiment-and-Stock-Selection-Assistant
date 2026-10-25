"""The Task 7 narrative layer: code writes the facts and numbers, the model rewrites the words
(mocked model, no network). Pinned to the stored-records sample of 17 Sep."""
import json
import sys
import unittest
from unittest.mock import patch

from sgx_capstone.configuration import ROOT, load_config, configure_report_modules

GOOD_CARDS = {
    "c1": "SGX Group dropped {c1_period_move} over the last five days, most of it in one session "
          "({c1_big_day_move}). Mostly positive news does not explain the fall. Task 6 expects a further fall "
          "(chance {c1_chance}), but its record here trails a simple rule, so it deserves little weight.",
    "c2": "UOB gained {c2_period_move} over the period. The news was mixed, so it does not clearly explain the "
          "move. Task 6 expects a rise (chance {c2_chance}), but its record on UOB trails a simple rule.",
    "c3": "Keppel Corp slipped {c3_period_move} with no single large day. Its news was mostly positive, so the "
          "news does not explain it. Task 6 expects no big move (chance {c3_chance}); its record trails a "
          "simple rule.",
}
SETTINGS = {"model": "openai/gpt-oss-20b", "endpoint": "https://example.invalid", "max_retries": 1}
DATE = "2026-09-17"


class NarrativeChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg, _ = load_config()
        sys.path.insert(0, str(ROOT / "src"))
        configure_report_modules(cls.cfg)
        from sgx_capstone.daily_report import gather_sample
        cls.D = gather_sample(ROOT, date=DATE)
        from sgx_capstone import narrative as N
        cls.N = N
        facts, _, _ = N.build_facts(cls.D)
        cls.GOOD = N.summary_draft(cls.D, facts)          # an unchanged draft is a valid rewrite

    def reply(self, summary=None, cards=None):
        return json.dumps({"summary": summary or self.GOOD, "cards": cards or GOOD_CARDS}), None

    def build(self, raw):
        with patch.object(self.N, "call_llm", return_value=raw):
            return self.N.build(self.D, "key", SETTINGS)

    def line(self, i, text):
        out = list(self.GOOD)
        out[i] = text
        return out

    def test_no_key_means_template(self):
        out = self.N.build(self.D, None, SETTINGS)
        self.assertEqual(out["mode"], "template")
        self.assertIsNone(out["summary"])

    def test_market_read_is_built_by_rules_from_tasks_1_to_6(self):
        lines = [t for _, t in self.N.market_read(self.D)]
        self.assertEqual(len(lines), 6)
        self.assertIn("Today the 15 stocks moved +0.3% on average", lines[0])
        self.assertIn("Financial Services (-9.8%)", lines[0])
        self.assertIn("For SGX Group and Keppel Corp, the company news and the price disagree", lines[1])
        self.assertIn("VIX 16.0", lines[2])
        self.assertIn("SGX Group (not yet in the price)", lines[3])
        self.assertIn("none with a ✓ record", lines[4])
        self.assertIn("have the highest chance of no big move", lines[5])
        self.assertNotIn("{", " ".join(lines))

    def test_good_draft_is_used_and_numbers_come_from_code(self):
        out = self.build(self.reply())
        self.assertEqual(out["mode"], "llm", out["flags"])
        self.assertEqual(out["flags"], [])
        self.assertIn("Today the 15 stocks moved +0.3% on average", out["summary"][0])
        self.assertEqual(len(out["cards"]), 3)
        self.assertNotIn("{", " ".join(out["summary"]) + " ".join(out["cards"].values()))
        self.assertIn("dropped 9.8%", out["cards"]["S68.SI"])

    def test_each_rule_replaces_only_the_failing_block(self):
        bad = {
            "N3 digit": (self.line(0, "Most stocks edged up: 10 rose."), None, "summary"),
            "N2 advice": (None, {**GOOD_CARDS, "c1": "Investors should buy SGX Group after a fall of "
                                                      "{c1_period_move}."}, "c1"),
            "N3 unknown": (None, {**GOOD_CARDS, "c2": "UOB moved {c2_target}."}, "c2"),
            "N4 other card": (None, {**GOOD_CARDS, "c2": "UOB moved like SGX Group ({c1_period_move})."}, "c2"),
            "N5 direction": (None, {**GOOD_CARDS, "c1": "SGX Group rose {c1_period_move} this week."}, "c1"),
            "N6 jargon": (None, {**GOOD_CARDS, "c3": "Keppel Corp looks bearish after {c3_period_move}."}, "c3"),
        }
        for rule, (summary, cards, block) in bad.items():
            out = self.build(self.reply(summary, cards))
            self.assertTrue(any(f.startswith(block) and rule.split()[0] in f for f in out["flags"]), (rule, out))
            if block == "summary":
                self.assertIsNone(out["summary"])
                self.assertEqual(len(out["cards"]), 3)
            else:
                self.assertIsNotNone(out["summary"])
                self.assertEqual(len(out["cards"]), 2)

    def test_reply_shapes_and_design_constants(self):
        top_level = json.dumps({"summary": self.GOOD, **GOOD_CARDS}), None
        self.assertEqual(len(self.build(top_level)["cards"]), 3)
        as_list = json.dumps({"summary": self.GOOD, "cards": list(GOOD_CARDS.values())}), None
        self.assertEqual(len(self.build(as_list)["cards"]), 3)
        constants = self.line(0, "Across the 15 stocks, {n_up} rose and {n_down} fell over the day; over the "
                                 "3 trading days ahead nothing is assumed.")
        out = self.build(self.reply(constants))
        self.assertIsNotNone(out["summary"], out["flags"])
        self.assertIn("raw", out)

    def test_live_run_failure_modes(self):
        one_string = json.dumps({"summary": [" ".join(self.GOOD)], "cards": GOOD_CARDS}), None
        out = self.build(one_string)
        self.assertEqual(len(out["summary"]), 1, out["flags"])
        compare = {**GOOD_CARDS, "c1": "SGX Group moved {c1_today_move} today, matching its {c1_period_move} "
                                       "over {period}. Task 6 expects a fall."}
        self.assertTrue(any("N8" in f for f in self.build(self.reply(None, compare))["flags"]))
        explains = {**GOOD_CARDS, "c2": "UOB gained {c2_period_move}. The mixed news explains the rise. Task 6 "
                                        "expects a rise, but its record trails a simple rule."}
        self.assertTrue(any("N9" in f for f in self.build(self.reply(None, explains))["flags"]))
        claims = self.line(1, self.GOOD[1].replace("the company news and the price disagree",
                                                   "the news explains the fall"))
        self.assertTrue(any("summary: N9" in f for f in self.build(self.reply(claims))["flags"]))

    def test_card_placeholders_sign_and_weight(self):
        card_ph = self.line(3, self.GOOD[3][:-1] + ", chance {c1_chance}.")
        self.assertTrue(any(f.startswith("summary: N4") for f in self.build(self.reply(card_ph))["flags"]))
        over = {**GOOD_CARDS, "c1": "SGX Group fell over {c1_period_move}. Mostly positive news does not explain "
                                    "the fall. Task 6 trails a simple rule, so it deserves little weight."}
        out = self.build(self.reply(None, over))
        self.assertIn("fell 9.8% over 11–17 Sep", out["cards"]["S68.SI"])
        weight = {**GOOD_CARDS, "c2": "UOB gained {c2_period_move}. The mixed news does not clearly explain it. "
                                      "Task 6's record gives its forecast moderate weight."}
        self.assertTrue(any("N10" in f for f in self.build(self.reply(None, weight))["flags"]))

    def test_rewrite_keeps_to_the_facts_of_its_draft_line(self):
        added = self.line(3, self.GOOD[3][:-1] + " and for UOB.")
        self.assertTrue(any("N11 adds 'uob'" in f for f in self.build(self.reply(added))["flags"]))
        dropped = self.line(3, "Today's news added no signal for any stock.")
        self.assertTrue(any("N11 drops 'sgx group'" in f for f in self.build(self.reply(dropped))["flags"]))

    def test_lessons_from_the_24_oct_live_run(self):
        """The model wrote chances in words and said the price had already followed when the fact
        said not yet; N13 and N12 catch both, and revalidate() drops such saved blocks."""
        ok = self.N.check_block("UOB rose {c2_period_move}.", {"c2_period_move": ("+1%", "")}, own="c2",
                                price_fact="already in the price")
        self.assertEqual(ok, [])
        fl = self.N.check_block("Task 6 predicts no big move with a fifty-percent chance.", {}, own="c1")
        self.assertTrue(any(f.startswith("N13") for f in fl), fl)
        fl = self.N.check_block("The price already reflects the news.", {}, own="c1",
                                price_fact="not yet in the price")
        self.assertTrue(any(f.startswith("N12") for f in fl), fl)
        fl = self.N.check_block("The news is not yet in the price.", {}, own="c1",
                                price_fact="the price went the other way")
        self.assertTrue(any(f.startswith("N12") for f in fl), fl)
        good = self.build(self.reply())
        saved = dict(good, cards=dict(good["cards"], **{"S68.SI": "The price already reflects the news, "
                                                                  "a fifty-percent chance."}))
        out = self.N.revalidate(saved, self.D)
        self.assertNotIn("S68.SI", out["cards"])
        self.assertIsNotNone(out["summary"])
        self.assertTrue(any("saved card S68.SI" in f for f in out["flags"]))

    def test_corporate_buyout_is_not_advice(self):
        from sgx_capstone.daily_report import summary_is_clean, clean_t5_summary
        self.assertTrue(summary_is_clean("Recent news on the CityCab buy‑out does not shift the stance."))
        self.assertFalse(summary_is_clean("Investors may buy the stock on weakness."))
        self.assertIn("a negative lean", clean_t5_summary("The linear model signals bearishness here."))
        full = clean_t5_summary("A positive geopolitical backdrop supports it. The outlook stays neutral with a "
                                "confidence of 0.55. Investors should monitor VIX spikes.")
        self.assertEqual(full, "A positive geopolitical backdrop supports it.")

    def test_retry_with_feedback_and_separate_calls(self):
        bad = json.dumps({"summary": self.line(0, "Normal 3‑day swings are wide: 10 rose.")})
        good = json.dumps({"summary": self.GOOD})
        cards_only = json.dumps({"cards": GOOD_CARDS})
        with patch.object(self.N, "call_llm", side_effect=[(bad, None), (good, None), (cards_only, None)]) as m:
            out = self.N.build(self.D, "key", SETTINGS)
        self.assertEqual(out["retried"], ["summary"], out["flags"])
        self.assertEqual(len(out["cards"]), 3)
        self.assertEqual(m.call_count, 3)
        self.assertIn("broke these checks", m.call_args_list[1].args[0][-1]["content"])
        hyphen = json.dumps({"summary": self.line(4, self.GOOD[4].replace("3-day", "3‑day"))})
        with patch.object(self.N, "call_llm", side_effect=[(hyphen, None), (cards_only, None)]):
            out = self.N.build(self.D, "key", SETTINGS)
        self.assertIsNotNone(out["summary"], out["flags"])

    def test_unusable_reply_falls_back_entirely(self):
        out = self.build(("not json", None))
        self.assertEqual(out["mode"], "template")
        self.assertIn("N1", out["flags"][0])
        out = self.build((None, "Timeout"))
        self.assertEqual(out["mode"], "template")

    def test_saved_text_for_an_older_layout_is_not_reused(self):
        old = {"mode": "llm", "summary": ["one", "two", "three", "four"], "cards": {"C52.SI": "x"}}
        self.assertFalse(self.N.matches_layout(old, self.D))
        self.assertTrue(self.N.matches_layout(self.build(self.reply()), self.D))

    def test_report_prints_the_ai_lines_and_labels_them(self):
        from sgx_capstone.daily_report import render
        from docx import Document
        import tempfile
        D = dict(self.D)
        D["narrative"] = self.build(self.reply())
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document(render(D, tmp))
            body = "\n".join([p.text for p in doc.paragraphs] +
                             [c.text for t in doc.tables for r in t.rows for c in r.cells]).lower()
        self.assertIn("the ai writer (openai/gpt-oss-20b, the same model that writes the news summaries)", body)
        # the reader never meets the internal stage names
        self.assertNotRegex(body, r"\btask [5-7]\b")
        self.assertIn("in short (ai-written, checked)", body)
        import task7_generate as g
        for pattern in g._R1_PATTERNS + g._R2_PATTERNS + g._R3_PATTERNS:
            self.assertNotRegex(body, pattern)


if __name__ == "__main__":
    unittest.main()
