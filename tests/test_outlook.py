"""Offline checks for the next-three-session outlook and the Daily Signal Report."""
import json
import pickle
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd

from sgx_capstone.configuration import ROOT, load_config, configure_report_modules
from sgx_capstone import outlook as O


class OutlookChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg, _ = load_config()
        sys.path.insert(0, str(ROOT / "src"))
        configure_report_modules(cls.cfg)
        with open(cls.cfg["paths"]["t6_artifacts"], "rb") as fh:
            cls.art = pickle.load(fh)
        rows = pd.read_csv(ROOT / "tests" / "fixtures" / "feature_rows.csv")
        cls.X = rows[rows.date == "2026-09-18"].set_index("ticker")[cls.art["feature_names"]]
        cls.ev = O.load_evidence(ROOT)

    def test_group_shapley_adds_up_exactly(self):
        shap = O.group_shapley(self.art, "RandomF", self.X)
        for t, s in shap.items():
            self.assertAlmostEqual(sum(s["move_by_group"].values()),
                                   s["p_move"] - s["p_move_average_day"], places=9)
            self.assertAlmostEqual(sum(s["lean_by_group"].values()),
                                   s["lean"] - s["lean_average_day"], places=9)

    def test_every_feature_belongs_to_one_group(self):
        grouped = sum(O.FEATURE_GROUPS.values(), [])
        self.assertEqual(sorted(grouped), sorted(self.art["feature_names"]))
        self.assertEqual(len(grouped), len(set(grouped)))
        self.assertEqual(set(O.PLAIN), set(self.art["feature_names"]))

    def test_missing_inputs_take_training_mean_not_zero(self):
        row = self.X.iloc[0].copy()
        row["vix"] = np.nan
        row["rsi"] = np.nan
        filled, missing = O.fill_missing(row, self.art)
        mean = O.average_day(self.art)
        self.assertEqual(sorted(missing), ["rsi", "vix"])
        self.assertAlmostEqual(filled["vix"], mean["vix"])
        self.assertGreater(filled["rsi"], 30)
        src = (ROOT / "sgx_capstone" / "legacy" / "task6_runtime.py").read_text()
        self.assertIn("SCALER_T6.mean_", src)
        self.assertNotIn("v = 0\n", src)

    def test_typical_range_is_sd_times_root_three(self):
        r = pd.Series([1.0, -1.0] * 10)
        out = O.typical_range(r)
        self.assertAlmostEqual(out["half_width_pct"], r.std(ddof=1) * np.sqrt(3))
        self.assertIsNone(O.typical_range(pd.Series([0.5, -0.5])))

    def test_outlook_calls_five_quieter_and_never_a_direction(self):
        returns = pd.DataFrame(np.random.default_rng(0).normal(0, 1, (25, 15)),
                               columns=self.X.index)
        looks = O.build_outlook(self.art, "RandomF", self.X, returns, self.ev)
        self.assertEqual(sum(o["activity"] == "quieter than most" for o in looks.values()), 5)
        for o in looks.values():
            text = O.render_explanation(o, self.ev)
            ok, flags = O.validate_explanation(text, O.explanation_facts(o, self.ev))
            self.assertTrue(ok, flags)
            self.assertNotRegex(text.lower(), r"\b(bullish|bearish|will rise|will fall)\b")

    def test_validator_rejects_advice_direction_and_invented_numbers(self):
        returns = pd.DataFrame(np.random.default_rng(1).normal(0, 1, (25, 15)),
                               columns=self.X.index)
        o = O.build_outlook(self.art, "RandomF", self.X, returns, self.ev)["ME8U.SI"]
        facts = O.explanation_facts(o, self.ev)
        lead = facts["group_effects_points"][0]["group"]
        bad = {
            "advice": f"{lead} lowered the chance. You should buy now.",
            "direction": f"{lead} lowered the chance, and the price is likely to rise.",
            "number": f"{lead} lowered the chance to 12.5%.",
        }
        for kind, text in bad.items():
            ok, flags = O.validate_explanation(text, facts)
            self.assertFalse(ok, kind)
        ok, flags = O.validate_explanation("Market mood lowered it.", facts)
        self.assertTrue(any(f.startswith("O6") for f in flags))

    def test_explanation_is_template_only(self):
        returns = pd.DataFrame(np.ones((25, 15)), columns=self.X.index)
        o = O.build_outlook(self.art, "RandomF", self.X, returns, self.ev)["D05.SI"]
        e = O.explain(o, self.ev)
        self.assertEqual(e["mode"], "template")
        self.assertEqual(e["text"], O.render_explanation(o, self.ev))
        src = (ROOT / "sgx_capstone" / "outlook.py").read_text()
        self.assertNotIn("requests.post", src)

    def test_checks_cannot_see_what_a_number_is_attached_to(self):
        # The limit that led to templated numbers: a sentence that reverses the
        # largest driver's effect and quotes the result as the driver's number
        # still passes every check, because each number exists in the facts.
        returns = pd.DataFrame(np.random.default_rng(1).normal(0, 1, (25, 15)),
                               columns=self.X.index)
        o = O.build_outlook(self.art, "RandomF", self.X, returns, self.ev)["ME8U.SI"]
        facts = O.explanation_facts(o, self.ev)
        lead = facts["group_effects_points"][0]
        wrong = "raised" if lead["points"] < 0 else "lowered"
        text = f"{lead['group'].capitalize()} {wrong} the chance of a move to {facts['chance_of_move_pct']}%."
        ok, flags = O.validate_explanation(text, facts)
        self.assertTrue(ok, flags)

    def test_r14_evidence_matches_retained_predictions(self):
        from sgx_capstone.research import r14_outlook
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            (tmp / "outputs" / "research").mkdir(parents=True)
            (tmp / "evidence").mkdir()
            src = ROOT / "outputs" / "research" / "walk_forward_predictions.csv"
            (tmp / "outputs" / "research" / src.name).write_bytes(src.read_bytes())
            fresh = r14_outlook.run(tmp)
        for key in ("quiet_move_rate", "rest_move_rate", "dates_quiet_strictly_lower", "dates_tied"):
            self.assertAlmostEqual(fresh["quiet"][key], self.ev["quiet"][key])
        self.assertAlmostEqual(fresh["range"]["range_coverage"], self.ev["range"]["range_coverage"])
        self.assertLess(self.ev["quiet"]["quiet_move_rate"], self.ev["quiet"]["rest_move_rate"])
        self.assertTrue(self.ev["by_model"]["RandomF"]["quiet_record"])
        self.assertFalse(self.ev["by_model"]["LogReg"]["quiet_record"])

    def test_stored_daily_returns_compound_to_every_stored_multi_day_return(self):
        from sgx_capstone.daily_report import stored_daily_returns
        from sgx_capstone.research.r14_outlook import SESSION
        daily, filled = stored_daily_returns(ROOT)
        wf = pd.read_csv(ROOT / "outputs" / "research" / "walk_forward_predictions.csv")
        wf = wf[(wf.calib == "oof") & (wf.model == "RandomF")]
        checked = 0
        for h in (3, 5):
            for d, g in wf[wf.h == h].groupby("date"):
                sess = [(pd.Timestamp(d) + k * SESSION).strftime("%Y-%m-%d") for k in range(1, h + 1)]
                if all(s in daily.index for s in sess):
                    comp = ((1 + daily.loc[sess] / 100).prod() - 1) * 100
                    self.assertLess(float((comp - g.set_index("ticker").ret).abs().max()), 1e-9)
                    checked += 1
        self.assertGreater(checked, 40)
        self.assertIn("2026-09-17", daily.index)

    def test_sample_report_builds_without_database_and_without_hindsight(self):
        from sgx_capstone.daily_report import build_sample
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path = build_sample(out_dir=tmp, date="2026-09-17")
            doc = Document(path)
            body = "\n".join([p.text for p in doc.paragraphs] +
                              [c.text for t in doc.tables for r in t.rows for c in r.cells]).lower()
            rec = json.loads(next(Path(tmp).glob("T7_report_input_SAMPLE_*.json")).read_text())
        self.assertIn("sample built from stored project records", body)
        for heading in ("today in 30 seconds", "which parts can you rely on?", "all 15 stocks today",
                        "the last 5 trading days and what may explain them", "the next 3 trading days",
                        "what these numbers are worth", "the system behind this report",
                        "what three months of testing show", "limitations",
                        "words used in this report"):
            self.assertIn(heading, body)
        # The report is for investors; how it was produced is saved next to it.
        self.assertNotIn("for reviewers", body)
        # No advice, targets or certainty anywhere in the report, tables included (checks R1-R3).
        import task7_generate as g
        for pattern in g._R1_PATTERNS + g._R2_PATTERNS + g._R3_PATTERNS:
            self.assertNotRegex(body, pattern)
        self.assertEqual(len(rec["rows"]), 15)
        self.assertEqual(sum(o["activity"] == "quieter than most" for o in rec["outlook"].values()), 5)
        # Every forecast in the quoted records had matured by the report date:
        # the champion record (R10) and the outlook record (R14 as of the date).
        from sgx_capstone.research.r14_outlook import SESSION
        last = pd.Timestamp(rec["record"]["last_forecast"]) + 3 * SESSION
        self.assertLessEqual(last, pd.Timestamp(rec["date"]))
        self.assertEqual(rec["record"]["as_of"], rec["date"])
        out_last = pd.Timestamp(rec["outlook_record"]["window"].split(" to ")[-1]) + 3 * SESSION
        self.assertLessEqual(out_last, pd.Timestamp(rec["date"]))
        # Stock cards name the period's dates and list the fired Task 5 rules,
        # re-run offline only where they reproduce the stored rl_v8 score.
        self.assertEqual(rec["date"], "2026-09-17")
        self.assertIn("over 11–17 sep", body)
        for label in ("news today", "price chart", "news & signals reading", "rules behind the reading", "(pattern:",
                      "today: news vs price", "last 5 days in context", "forecast lean", "record vs simple rule",
                      "champion", "why.", "backdrop.", "for you."):
            self.assertIn(label, body)
        # Task 6: the stored live forecast for the day is shown for every stock, with its record.
        self.assertTrue(all(r["t6_source"] == "stored live forecast" for r in rec["rows"].values()))
        self.assertEqual(sum(w in body for w in ("no big move", "rise", "fall")), 3)
        self.assertGreater(rec["findings"]["n"], 200)
        # The context line draws its conclusion from news, sector peers and the market, never from chart signals.
        self.assertIn("so the company news and the price disagree", body)
        # Reliability is stated next to each kind of statement.
        for verdict in ("reliable", "promising", "not reliable", "describes only"):
            self.assertIn(verdict, body)
        detail = pd.read_csv(ROOT / "outputs" / "research" / "R8_rl_v8_rule_detail.csv")
        days = detail[detail.date == rec["date"]].drop_duplicates("ticker")
        self.assertEqual(int(days.reproduced.sum()), 15)
        for t, r in rec["rows"].items():
            stored = days.set_index("ticker").loc[t, "net_stored"]
            self.assertAlmostEqual(r["net"], stored, places=3)
        # The political-tone rule reads FinBERT confidence, not tone: never shown as a signal.
        self.assertNotIn("geopolitical", body)
        from sgx_capstone.daily_report import strip_political_sentences
        self.assertEqual(strip_political_sentences(
            "A supportive geopolitical backdrop. However, results beat estimates."),
            "Results beat estimates.")


    def test_default_sample_is_the_saved_21_sep_database_copy(self):
        from sgx_capstone.daily_report import build_sample
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path = build_sample(out_dir=tmp)
            doc = Document(path)
            body = "\n".join([p.text for p in doc.paragraphs] +
                              [c.text for t in doc.tables for r in t.rows for c in r.cells]).lower()
            rec = json.loads(next(Path(tmp).glob("T7_report_input_SAMPLE_*.json")).read_text())
            notes = next(Path(tmp).glob("T7_reviewer_notes_SAMPLE_*.md")).read_text()
        self.assertEqual(rec["date"], "2026-09-21")
        # Cards show where today's news adds something: the three positive readings, strongest first.
        self.assertEqual(rec["cards"], ["S68.SI", "U11.SI", "A17U.SI"])
        self.assertIn("where today's news adds something", body)
        # 'Calmer than most' comes from the forecast stored that day: no hindsight, no training-period caveat.
        self.assertEqual(rec["calm_source"], "stored")
        snap = json.loads((ROOT / "tests" / "fixtures" / "live_snapshot_2026-09-21.json").read_text())["D"]["rows"]
        top5 = sorted(snap, key=lambda t: -snap[t]["t6_probs"]["neutral"])[:5]
        self.assertEqual({t for t, o in rec["outlook"].items() if o["activity"] == "quieter than most"}, set(top5))
        self.assertIn("rebuilt offline from the saved copy", body)
        self.assertNotIn("sample rebuilt", body)
        self.assertNotIn("training period", body)
        # Task 5's AI summary in full, advice sentences removed; the political rule flagged once.
        self.assertIn("remains positive on the back of a strong rsi oversold signal", body)
        self.assertNotIn("investors should monitor", body)
        self.assertIn("those two rules fire every day and cancel out", body)
        self.assertIn("borderline: without the two political-news rules", body)
        self.assertNotIn("for reviewers", body)
        self.assertIn("## Provenance", notes)
        import task7_generate as g
        for pattern in g._R1_PATTERNS + g._R2_PATTERNS + g._R3_PATTERNS:
            self.assertNotRegex(body, pattern)


if __name__ == "__main__":
    unittest.main()
