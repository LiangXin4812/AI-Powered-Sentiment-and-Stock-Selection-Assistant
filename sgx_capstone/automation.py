"""Daily automation and the multi-date check (20 Oct 2026).

run-all       after the SGX close: Tasks 1-4 for today -> Task 5 and Task 6 ->
              the Daily Signal Report -> reports/<date>/ (docx, PDF when
              LibreOffice is installed, reviewer notes) and reports/README.md,
              the index a reviewer opens in the browser; optional e-mail.
check-dates   builds the report for every stored date in a range from the live
              database and checks each one (no AI call unless --llm), so a
              reviewer can see the report is stable beyond the sample date.

Both read credentials from the environment only (SUPABASE_URL, GROQ_API_KEY,
optional SMTP_HOST / SMTP_PORT / SMTP_USER / SMTP_PASSWORD / REPORT_TO).
"""
from __future__ import annotations

import csv
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

SGT = ZoneInfo("Asia/Singapore")
SECTION_HEADINGS = ("today in 30 seconds", "all 15 stocks today",
                    "the last 5 trading days and what may explain them", "the next 3 trading days",
                    "what these numbers are worth", "the system behind this report",
                    "what three months of testing show", "limitations", "words used in this report")


# ── trading day ────────────────────────────────────────────────────────────

def is_trading_day(date: str, probe: str = "D05.SI") -> bool:
    """True when SGX printed a daily bar for `date` (weekends and public
    holidays have none), so no holiday table has to be maintained."""
    import yfinance as yf
    d = datetime.strptime(date, "%Y-%m-%d").date()
    if d.weekday() >= 5:
        return False
    bars = yf.Ticker(probe).history(period="7d", interval="1d", auto_adjust=False)
    return any(ts.date() == d for ts in bars.index)


# ── publishing ─────────────────────────────────────────────────────────────

def to_pdf(docx: Path) -> Path | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=False, capture_output=True, timeout=300)
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def publish(root: Path, docx: Path, date: str) -> dict:
    """Copy the day's report into reports/<date>/ and rebuild reports/README.md."""
    dest = Path(root) / "reports" / date
    dest.mkdir(parents=True, exist_ok=True)
    out = {"docx": dest / docx.name}
    shutil.copy2(docx, out["docx"])
    notes = docx.parent / docx.name.replace("T7_Daily_Signal_Report_", "T7_reviewer_notes_").replace(".docx", ".md")
    if notes.exists():
        out["notes"] = dest / notes.name
        shutil.copy2(notes, out["notes"])
    pdf = to_pdf(out["docx"])
    if pdf:
        out["pdf"] = pdf
    write_index(Path(root) / "reports")
    return out


def write_index(reports: Path) -> Path:
    days = sorted((p for p in reports.iterdir() if p.is_dir() and re.fullmatch(r"\d{4}-\d{2}-\d{2}", p.name)),
                  reverse=True)
    lines = ["# SGX Daily Signal Reports", "",
             "Generated automatically after the SGX close on each trading day (Tasks 1-7). "
             "General information, not financial advice.", "",
             "| Date | Report | PDF | How it was produced |", "|---|---|---|---|"]
    for d in days:
        docx = next((p.name for p in d.glob("*.docx")), None)
        pdf = next((p.name for p in d.glob("*.pdf")), None)
        md = next((p.name for p in d.glob("*.md")), None)
        cell = lambda name: f"[{name.split('.')[-1]}]({d.name}/{name})" if name else "—"
        lines.append(f"| {d.name} | {cell(docx)} | {cell(pdf)} | {cell(md)} |")
    path = reports / "README.md"
    path.write_text("\n".join(lines) + "\n")
    return path


def email(files: dict, date: str) -> bool:
    """Send the report when SMTP settings are present; silently skip otherwise."""
    host, to = os.environ.get("SMTP_HOST"), os.environ.get("REPORT_TO")
    user, pwd = os.environ.get("SMTP_USER"), os.environ.get("SMTP_PASSWORD")
    if not (host and to and user and pwd):
        return False
    import smtplib
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["Subject"] = f"SGX Daily Signal Report {date}"
    msg["From"], msg["To"] = user, to
    msg.set_content(f"The SGX Daily Signal Report for {date} is attached.\n\n"
                    "General information, not financial advice.")
    for key in ("pdf", "docx"):
        if key in files:
            data = Path(files[key]).read_bytes()
            sub = "pdf" if key == "pdf" else "vnd.openxmlformats-officedocument.wordprocessingml.document"
            msg.add_attachment(data, maintype="application", subtype=sub, filename=Path(files[key]).name)
    with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "587")), timeout=60) as s:
        s.starttls()
        s.login(user, pwd)
        s.send_message(msg)
    return True


# ── checking a report ──────────────────────────────────────────────────────

def report_text(docx: Path) -> str:
    from docx import Document
    doc = Document(str(docx))
    return "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])


def check_report(docx: Path, D: dict) -> list[str]:
    """Problems found in one built report (empty list = passes)."""
    import task7_generate as g
    from .daily_report import pick_cards, _calm
    from . import narrative as N
    body = report_text(docx)
    low = body.lower()
    issues = []
    for p in g._R1_PATTERNS + g._R2_PATTERNS + g._R3_PATTERNS:
        m = re.search(p, low)
        if m:
            issues.append(f"compliance: '{m.group(0)}'")
    issues += [f"missing section '{h}'" for h in SECTION_HEADINGS if h not in low]
    if "for reviewers" in low:
        issues.append("reviewer page printed")
    if re.search(r"\(page \d", low):
        issues.append("page-number reference (sections are numbered instead)")
    if "{" in body:
        issues.append("unfilled placeholder")
    if re.search(r"\bnan\b", body):
        issues.append("'nan' printed")
    if len(D["rows"]) != 15:
        issues.append(f"{len(D['rows'])} stocks")
    if len(_calm(D)) != 5:
        issues.append(f"{len(_calm(D))} calmer stocks")
    if len(pick_cards(D)) != 3:
        issues.append(f"{len(pick_cards(D))} cards")
    if len(N.market_read(D)) != 6:
        issues.append("market read is not six lines")
    return issues


def check_dates(ns, cfg, start: str, end: str, root: Path, api_key=None) -> Path:
    """Build and check the report for every stored date in [start, end]."""
    from sqlalchemy import text
    from .daily_report import gather_live, render, pick_cards, short
    from . import narrative as N
    with ns["SUPABASE_ENGINE"].connect() as conn:
        dates = [str(r[0]) for r in conn.execute(text("""
            SELECT p.prediction_date FROM task6_predictions p
            WHERE p.is_champion = TRUE AND p.prediction_date BETWEEN CAST(:a AS DATE) AND CAST(:b AS DATE)
            GROUP BY p.prediction_date HAVING COUNT(DISTINCT p.ticker) = 15
            INTERSECT
            SELECT CAST(analysis_date AS DATE) FROM llm_analyses
            WHERE CAST(analysis_date AS DATE) BETWEEN CAST(:a AS DATE) AND CAST(:b AS DATE)
            GROUP BY CAST(analysis_date AS DATE) HAVING COUNT(DISTINCT ticker) = 15
            ORDER BY 1"""), {"a": start, "b": end})]
    out_root = Path(root) / "outputs" / "date_check"
    out_root.mkdir(parents=True, exist_ok=True)
    rows = []
    for d in dates:
        row = {"date": d, "ok": False, "issues": "", "cards": "", "positive": "", "negative": "",
               "flagged": "", "calm_source": "", "narrative": "", "error": ""}
        try:
            D = gather_live(ns, cfg, d)
            D["narrative"] = N.build(D, api_key, cfg.get("llm"))
            path = render(D, out_root / d)
            issues = check_report(path, D)
            row.update({"ok": not issues, "issues": "; ".join(issues),
                        "cards": ", ".join(short(r["name"]) for r in pick_cards(D)),
                        "positive": sum(r["sentiment"] == "bullish" for r in D["rows"].values()),
                        "negative": sum(r["sentiment"] == "bearish" for r in D["rows"].values()),
                        "flagged": sum(r["attention"] in ("ELEVATED", "NOTABLE") for r in D["rows"].values()),
                        "calm_source": D.get("calm_source"), "narrative": D["narrative"]["mode"]
                        + (f" ({len(D['narrative']['flags'])} block(s) replaced)" if D["narrative"]["flags"] else "")})
        except Exception as e:                                   # noqa: BLE001
            row["error"] = f"{type(e).__name__}: {e}"[:300]
        rows.append(row)
        print(f"  {d}: {'OK' if row['ok'] else 'PROBLEM'}  cards: {row['cards'] or '-'}"
              + (f"  issues: {row['issues']}" if row["issues"] else "") + (f"  error: {row['error']}" if row["error"] else ""))
    summary = out_root / "summary.csv"
    with summary.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]) if rows else ["date"])
        w.writeheader()
        w.writerows(rows)
    ok = sum(r["ok"] for r in rows)
    print(f"Checked {len(rows)} stored dates from {start} to {end}: {ok} passed every check. Summary: {summary}")
    return summary
